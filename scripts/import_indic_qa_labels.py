"""Turn indic-qa reviewer workbooks into Bridge human labels (pa, ml, hi, or).

indic-qa's review rounds are workbooks with one example sheet per rule group:
`example_id, rule_id, ref, context, finding, suggestion, …, verdict,
reviewer_correction, note`, and verdicts Agree / Agree with change /
Disagree. This writes them as a Bridge review round, so the existing human
gate (scripts/language_qa_benchmark.py) measures each indic-qa rule's
precision on Bridge's own findings and can promote it to inline:

    python scripts/import_indic_qa_labels.py --lang hi --irv-dir "D:\\Claude Lab\\IRV Hindi" \\
        --round 2026-10-03-final --workbook "<indic-qa>/outputs/hi-irv/hi-irv_qa-app_final_workbook_completed.xlsx"
    python scripts/language_qa_benchmark.py --human-labels benchmark/human/hi --language hin --gate

Output: benchmark/human/<code>/<round>/human_labels.jsonl and verses.jsonl
beside it. A language's rounds sit one level below benchmark/human/ so a book
id never collides with Tamil's (load_human_rounds keys verses by book).

Every label is anchored in the verse as Bridge's importer stores it, never in
indic-qa's own line offsets. A row is anchored only when its finding text
occurs in the verse exactly once, or once inside its context. It is then moved
onto the finding Bridge itself reports there for that rule, scanning exactly
the verses written to verses.jsonl (what the gate will scan). This mirrors the
Tamil rounds, which were labelled on the engine's own findings. A row in a
note, under a rule Bridge has off, or with no Bridge finding at that place is
counted and reported by reason, never guessed and never written. The verdicts
map to Bridge's codes:
Agree -> TP, Agree with change -> TP_OTHER_FIX (the finding is right, the fix
is not), Disagree -> FP. Rows that judge something else (word lists, open
questions) have no example_id/finding/verdict columns and are skipped.

Labels are data from a reviewer: this script never edits them, and a rule is
drawn inline only by a reviewed rule_versions.json change with a DECISIONS line.
"""
from __future__ import annotations

import argparse
import collections
import json
import sys
import unicodedata
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "engine"))

from tc_ai_bridge.language_packs import default_pack, indic_qa_vendor  # noqa: E402
from tc_ai_bridge.language_packs.indic_qa_adapter import book_key  # noqa: E402
from tc_ai_bridge.language_qa import lift_inline_usfm  # noqa: E402
from tc_ai_bridge.language_qa_benchmark import book_verses, scan_book  # noqa: E402

DECLARED = {"hi": "hin", "ml": "mal", "or": "ory", "pa": "pan"}
VERDICTS ={"agree": "TP", "agree with change": "TP_OTHER_FIX", "disagree": "FP"}
NEEDED = ("example_id", "rule_id", "ref", "finding", "verdict")
ELLIPSIS = "…"


def nfc(text: str) -> str:
    return unicodedata.normalize("NFC", text or "")


def sheet_rows(path: Path):
    """(sheet, row dict) for every example row of every sheet that has one."""
    from openpyxl import load_workbook  # tooling only (engine dev extra)
    book = load_workbook(path, read_only=True, data_only=True)
    for sheet in book.worksheets:
        rows = sheet.iter_rows(values_only=True)
        header = None
        for row in rows:
            names = [str(c).strip().lower() if c is not None else "" for c in row]
            if header is None:
                if all(n in names for n in NEEDED):
                    header = names
                continue
            values = dict(zip(header, row))
            if values.get("example_id"):
                yield sheet.title, values


def verse_key(chapters: dict[str, dict[str, str]], chapter: str, verse: str) -> str | None:
    """The stored key holding verse `verse`: itself, or a bridge ("3-4") or
    segment ("3a") that starts there or spans it."""
    keys = chapters.get(chapter, {})
    if verse in keys:
        return verse
    number = int(verse)
    for key in keys:
        head, _, tail = key.partition("-")
        first = int("".join(c for c in head if c.isdigit()) or -1)
        last = int("".join(c for c in tail if c.isdigit()) or first)
        if first <= number <= last:
            return key
    return None


def anchor(text: str, finding: str, context: str) -> tuple[int, int] | None:
    """(start, end) of `finding` in `text` when it is unambiguous."""
    finding = nfc(finding).strip(ELLIPSIS).strip()
    text = nfc(text)
    if not finding:
        return None
    starts = [i for i in range(len(text)) if text.startswith(finding, i)]
    if len(starts) == 1:
        return starts[0], starts[0] + len(finding)
    if len(starts) > 1 and context:
        # The context is a window around the finding, cut with "…" at either end.
        window = nfc(context).strip(ELLIPSIS).strip()
        at = text.find(window) if window else -1
        inside = [s for s in starts if at >= 0 and at <= s and s + len(finding) <= at + len(window)]
        if len(inside) == 1:
            return inside[0], inside[0] + len(finding)
    return None


def main() -> int:
    for stream in (sys.stdout, sys.stderr):
        stream.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--lang", required=True, choices=indic_qa_vendor.PROFILES)
    parser.add_argument("--irv-dir", type=Path, required=True, help="the IRV folder the workbook was built from")
    parser.add_argument("--workbook", type=Path, action="append", required=True)
    parser.add_argument("--round", required=True, help="the round's folder name, e.g. 2026-10-03-final")
    parser.add_argument("--out", type=Path, default=REPO / "benchmark" / "human")
    args = parser.parse_args()

    pack = default_pack(f"{args.lang}-irv")
    profile = indic_qa_vendor.profile(args.lang)
    switch = getattr(profile, "SHAPE_RULE_SWITCH", {})
    grouped = f"{args.lang}.norm.encoding" if f"{args.lang}.norm.encoding" in profile.RULES else ""
    sfms ={p.name: p for p in args.irv_dir.iterdir() if p.suffix.lower() in (".sfm", ".usfm")}
    books: dict[str, dict[str, dict[str, str]]] = {}
    for sfm in sorted(sfms.values()):
        book, chapters = book_verses(sfm)
        books[book_key(book)] = chapters

    labels, used = [], collections.defaultdict(dict)
    skipped: collections.Counter = collections.Counter()
    for workbook in args.workbook:
        for sheet, row in sheet_rows(workbook):
            verdict = VERDICTS.get(str(row.get("verdict") or "").strip().lower())
            if verdict is None:
                skipped[f"verdict {row.get('verdict')!r}"] += 1
                continue
            rule = switch.get(str(row["rule_id"]).strip(), str(row["rule_id"]).strip())
            if pack.by_id(rule) is None and rule.startswith(f"{args.lang}.norm.") and grouped:
                # A part of the grouped encoding fix (ml.norm.chillu-zwj, ...):
                # Bridge reports it as that one catalogue rule.
                rule = grouped
            if pack.by_id(rule) is None:
                skipped["rule not in the pack"] += 1
                continue
            try:
                book_code, cv = str(row["ref"]).strip().split(" ", 1)
                chapter, verse = cv.split(":", 1)
                int(chapter), int(verse)
            except ValueError:
                skipped["unreadable ref"] += 1
                continue
            if chapter == "0" or verse == "0":
                skipped["outside verses (title, heading, intro)"] += 1
                continue
            chapters = books.get(book_key(book_code))
            key = verse_key(chapters or {}, chapter, verse)
            if key is None:
                skipped["verse not in the IRV folder"] += 1
                continue
            text = chapters[chapter][key]
            span = anchor(text, str(row.get("finding") or ""), str(row.get("context") or ""))
            if span is None:
                skipped["finding not found once in the verse"] += 1
                continue
            suggestion = str(row.get("suggestion") or "").strip()
            labels.append({
                "id": f"{args.lang}-{row['example_id']}", "kind": "flagged", "rule": f"{pack.name}/{rule}",
                "book": book_key(book_code), "ch": chapter, "v": key, "start": span[0], "end": span[1],
                "original": text[span[0]:span[1]],
                "suggestions": [s.strip() for s in suggestion.split(",") if s.strip()][:5],
                "verdict": verdict, "correctForm": str(row.get("reviewer_correction") or ""),
                "note": str(row.get("note") or ""), "sheet": sheet, "source": workbook.name,
            })
            used[book_key(book_code)].setdefault(chapter, {})[key] = text

    # A label credits the finding Bridge reports at exactly its span, as the
    # Tamil rounds were labelled on the engine's own findings. Scan the same
    # verses the gate will scan (verses.jsonl), then move each label onto the
    # overlapping finding of its rule. A row Bridge does not report is a
    # parity note for this script's output, not a label: written as one, the
    # gate would count it as a regression from the first run.
    found = {book: scan_book(book, chapters, language=DECLARED[args.lang])["findings"]
             for book, chapters in used.items()}
    kept, unmatched = [], collections.Counter()
    grouped_id = f"{pack.name}/{grouped}" if grouped else ""
    for label in labels:
        rule = pack.by_id(label["rule"].split("/", 1)[1])
        text = used[label["book"]][label["ch"]][label["v"]]
        lifted, _ = lift_inline_usfm(text)
        visible = set(lifted.raw_index) if lifted is not None else set()
        if not rule.enabled:
            skipped["rule off in Bridge"] += 1
            continue
        if not all(i in visible for i in range(label["start"], label["end"])):
            skipped["inside a note or markup (Bridge scans verse text only)"] += 1
            continue
        here = [f for f in found[label["book"]] if (f["chapter"], f["verse"]) == (label["ch"], label["v"])]
        overlap = [f for f in here if f["ruleId"] == label["rule"]
                   and f["start"] < label["end"] and label["start"] < f["end"]]
        if not overlap and grouped_id and label["rule"].startswith(f"{pack.name}/{args.lang}.norm."):
            # A part of the grouped encoding fix: Bridge reports the group.
            overlap = [f for f in here if f["ruleId"] == grouped_id
                       and f["start"] < label["end"] and label["start"] < f["end"]]
            if overlap:
                label = {**label, "rule": grouped_id}
        if not overlap:
            unmatched[label["rule"]] += 1
            continue
        best = min(overlap, key=lambda f: (abs(f["start"] - label["start"]) + abs(f["end"] - label["end"])))
        kept.append({**label, "start": best["start"], "end": best["end"], "original": best["originalText"],
                     "fid": best["id"]})
    # Two rows on one finding: keep the first (the gate scores a finding once).
    seen, labels = set(), []
    for label in kept:
        if label["fid"] not in seen:
            seen.add(label["fid"])
            labels.append(label)
    skipped["two rows on one Bridge finding"] += len(kept) - len(labels)
    # verses.jsonl keeps every verse scanned above, labelled or not: the gate
    # scans exactly these, so its word counts, and so its consistency
    # majorities, are the ones the labels were matched against.
    out = args.out / args.lang / args.round
    out.mkdir(parents=True, exist_ok=True)
    with (out / "human_labels.jsonl").open("w", encoding="utf-8", newline="\n") as handle:
        for label in labels:
            handle.write(json.dumps(label, ensure_ascii=False) + "\n")
    with (out / "verses.jsonl").open("w", encoding="utf-8", newline="\n") as handle:
        for book in sorted(used):
            for chapter in sorted(used[book], key=int):
                for verse, text in used[book][chapter].items():
                    handle.write(json.dumps({"book": book, "chapter": chapter, "verse": verse, "text": text},
                                            ensure_ascii=False) + "\n")
    print(f"{len(labels)} labels written to {out.relative_to(REPO)}")
    print("by verdict:", dict(collections.Counter(label["verdict"] for label in labels)))
    print("skipped:", dict(skipped))
    print("not reported by Bridge at that place, by rule:", dict(unmatched.most_common()))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
